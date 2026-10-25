"""Turn the HTML deck into a .pptx: one full-slide picture per page (as it looks when presented) + its speaker notes.

    python docs/deck/build.py        # first: the HTML deck
    python docs/deck/build_pptx.py   # then: docs/XFitness-Chatbot_Presentation-v<PPTX_VERSION>.pptx

Each page is shown in Chromium at 1920×1080, every animation run to its end, and captured at 2× (3840×2160);
the page's <aside> becomes the slide's notes. Needs playwright (with `playwright install chromium`) and python-pptx.
The pictures cannot be clicked: the playgrounds in the appendix work only in the HTML deck."""
import io
import sys
from pathlib import Path

from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.util import Emu

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from build import OUT, ROOT  # noqa: E402

PPTX_VERSION = "1.4"
PPTX = ROOT / f"docs/XFitness-Chatbot_Presentation-v{PPTX_VERSION}.pptx"
HIDE_UI = "#ui,#bar,#bartip,#menu,#notes,#zoom,.edge{display:none!important}"
SETTLE_MS = 2500     # the longest entrance animation (counters, bars, graph lines) is about 2 s


def capture() -> list[tuple[bytes, str]]:
    shots = []
    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page(viewport={"width": 1920, "height": 1080}, device_scale_factor=2)
        page.goto(OUT.as_uri())
        page.add_style_tag(content=HIDE_UI)
        page.evaluate("document.fonts.ready")
        total = page.evaluate("document.querySelectorAll('.slide').length")
        for n in range(total):
            page.evaluate(f"show({n})")
            page.wait_for_timeout(SETTLE_MS)
            notes = page.evaluate(f"(document.querySelectorAll('.slide')[{n}].querySelector('aside')||{{}}).innerText||''")
            shots.append((page.screenshot(type="jpeg", quality=90), notes.strip()))
            print(f"  page {n + 1}/{total}", end="\r")
        browser.close()
    return shots


def main() -> None:
    shots = capture()
    deck = Presentation()
    deck.slide_width, deck.slide_height = Emu(12192000), Emu(6858000)          # 16:9, as the earlier versions
    blank = deck.slide_layouts[6]
    for image, notes in shots:
        slide = deck.slides.add_slide(blank)
        slide.shapes.add_picture(io.BytesIO(image), 0, 0, deck.slide_width, deck.slide_height)
        slide.notes_slide.notes_text_frame.text = notes
    deck.save(PPTX)
    print(f"{PPTX.relative_to(ROOT)} · {len(shots)} slides · {PPTX.stat().st_size / 1e6:.1f} MB")


if __name__ == "__main__":
    main()
